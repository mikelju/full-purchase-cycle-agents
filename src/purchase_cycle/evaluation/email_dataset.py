"""Golden dataset of the email order evaluation: seeded plan, validation, rendering and build.

The seeded plan fixes, before any text exists, the category, sender, order or
not, and every line with its expected SKU (or out-of-catalog item), expected
quantity in sale units, unit expression, trap and location, so labels are
decided by construction. The coding agent writes subjects, bodies and line
texts; this module validates them and renders deterministic `.eml` files whose
attachments are built from the plan, so attachment contents match the labels.
"""

import csv
import html
import io
import random
import re
import zipfile
from collections import Counter, defaultdict
from datetime import UTC, datetime, timedelta
from difflib import SequenceMatcher
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from purchase_cycle.catalog import CUSTOMERS, PRODUCTS
from purchase_cycle.config import EVALS_DIR
from purchase_cycle.email_order import EmailRejected, model_text, parse_email
from purchase_cycle.evaluation.planning import (
    GENERIC_QUANTITIES,
    OUT_OF_CATALOG_ITEMS,
    WORD_QUANTITIES,
    read_jsonl,
    write_jsonl,
)
from purchase_cycle.evaluation.stats import wilson_interval
from purchase_cycle.quantities import has_number, pack_size
from purchase_cycle.web_form import match_key

DATASET_VERSION = "1.0"
SEED = 20261008
EMAILS_PER_CATEGORY = 52
DEV_PER_CATEGORY = 13

CATEGORIES = (
    "body_list",
    "body_conversational",
    "txt_attachment",
    "pdf_attachment",
    "xlsx_attachment",
    "not_an_order",
)
ORDER_CATEGORIES = CATEGORIES[:-1]
BODY_CATEGORIES = ("body_list", "body_conversational")
TABLE_CATEGORIES = ("pdf_attachment", "xlsx_attachment")
SOURCE = {
    "body_list": "body",
    "body_conversational": "body",
    "txt_attachment": "txt",
    "pdf_attachment": "pdf",
    "xlsx_attachment": "xlsx",
    "not_an_order": "body",
}
LAYOUTS = {
    "body_list": ("plain_body", "html_body", "plain_body", "plain_body"),
    "body_conversational": ("plain_body", "html_body", "plain_body", "plain_body"),
    "txt_attachment": ("text_list",),
    "pdf_attachment": ("order_form", "delivery_note"),
    "xlsx_attachment": ("header_row", "extra_columns"),
    "not_an_order": ("plain_body", "html_body", "plain_body", "plain_body"),
}
LINE_COUNTS = {
    "body_list": (1, 2, 2, 3, 3, 4, 4, 5, 6, 7, 8),
    "body_conversational": (1, 2, 2, 3, 3, 4),
    "txt_attachment": (1, 2, 2, 3, 3, 4, 4, 5, 6, 7, 8),
    "pdf_attachment": (1, 2, 2, 3, 3, 4, 4, 5, 6, 7, 8),
    "xlsx_attachment": (1, 2, 2, 3, 3, 4, 4, 5, 6, 7, 8),
}
# Fixed share of the lines of every order category; the rest are plain lines.
TRAP_SHARES = (("out_of_catalog", 0.10), ("typo", 0.10), ("near_miss", 0.10))
UNIT_WEIGHTS = {
    "written": (("sale_units", 7), ("quantity_words", 1), ("items_to_packs", 1), ("dozen", 1)),
    "conversational": (
        ("sale_units", 3),
        ("quantity_words", 2),
        ("items_to_packs", 2),
        ("dozen", 1),
        ("half_dozen", 1),
        ("couple", 1),
    ),
    "table": (("sale_units", 5), ("items_to_packs", 1)),
}
NOT_ORDER_KINDS = ("delivery_question", "complaint", "newsletter", "invoice_query", "product_question", "thank_you")
MAX_PRODUCT_SHARE = 0.01  # no product above 1% of the catalog lines
MAX_SUBJECT_CHARS = 120
MAX_LINE_CHARS = 160
MAX_TABLE_TEXT_CHARS = 80  # one PDF table cell, kept on one line
MAX_SHORT_BODY_CHARS = 400  # body of the attachment categories
MIN_CONTEXT_CHARS = 60  # conversational text around the lines
MAX_BODY_CHARS = 3000
BASE_DATE = datetime(2026, 9, 1, 8, 0, tzinfo=UTC)
FILE_DATE = (2026, 9, 1, 0, 0, 0)
SHOP_ADDRESS = "orders@purchase-cycle.example"

DATASET_DIR = EVALS_DIR / "datasets" / "email_order_extraction"
PLAN_PATH = DATASET_DIR / "plan.jsonl"
TEXTS_DIR = DATASET_DIR / "texts"
EMAILS_DIR = DATASET_DIR / "emails"
DATASET_PATH = DATASET_DIR / "dataset.jsonl"
REVIEW_PATH = DATASET_DIR / "second_pass_review.jsonl"
REVIEW_DECISIONS = ("agree", "label_fixed", "annotator_wrong")
MIN_LINE_SIMILARITY = 0.5
AUDIT_PATH = EVALS_DIR / "audit" / f"email_order_extraction-audit-v{DATASET_VERSION}.csv"
AUDIT_SEED = 40
AUDIT_SIZE = 40
MAX_WRONG = 1
AUDIT_COLUMNS = ("id", "category", "file", "email_text", "is_order", "expected_lines", "verdict", "comment")
VERDICTS = ("ok", "wrong")

PRODUCT_BY_SKU = {p.sku: p for p in PRODUCTS}
CUSTOMER_BY_CODE = {c.code: c for c in CUSTOMERS}


# ---------- planning ----------


def _unit(rng, weights_key: str, product) -> tuple[str, int, dict]:
    """Unit expression, expected quantity in sale units and extra trap fields."""
    pack = pack_size(product.sale_unit) if product else None
    styles = [(s, w) for s, w in UNIT_WEIGHTS[weights_key] if s != "items_to_packs" or (pack or 0) >= 10]
    if product is None:  # out-of-catalog lines keep the quantity simple
        styles = [(s, w) for s, w in styles if s in ("sale_units", "quantity_words")]
    style = rng.choices([s for s, _ in styles], weights=[w for _, w in styles])[0]
    if style == "quantity_words":
        return style, rng.choice(WORD_QUANTITIES), {}
    if style == "items_to_packs":
        quantity = rng.randint(1, 10)
        return style, quantity, {"items_requested": quantity * pack}
    if style == "dozen":
        return style, rng.choice([12, 24, 36, 48]), {}
    if style == "half_dozen":
        return style, 6, {}
    if style == "couple":
        return style, 2, {}
    return style, rng.choice(GENERIC_QUANTITIES), {}


def _closest_sibling(product, siblings) -> str:
    others = [s for s in siblings[product.family] if s.sku != product.sku]
    return max(others, key=lambda s: (SequenceMatcher(None, s.name, product.name).ratio(), s.sku)).sku


def _attachment_name(category: str, layout: str, number: int) -> str | None:
    return {
        "txt_attachment": f"order-{number:04d}.txt",
        "pdf_attachment": f"{'purchase-order' if layout == 'order_form' else 'delivery-note'}-{number:04d}.pdf",
        "xlsx_attachment": f"order-{number:04d}.xlsx",
    }.get(category)


def build_plan(seed: int = SEED) -> list[dict]:
    rng = random.Random(seed)
    products = sorted(PRODUCTS, key=lambda p: p.sku)
    siblings = defaultdict(list)
    for p in products:
        siblings[p.family].append(p)
    uses = Counter()
    emails = []
    for category in CATEGORIES:
        layouts = [LAYOUTS[category][i % len(LAYOUTS[category])] for i in range(EMAILS_PER_CATEGORY)]
        rng.shuffle(layouts)
        group = []
        if category == "not_an_order":
            kinds = [NOT_ORDER_KINDS[i % len(NOT_ORDER_KINDS)] for i in range(EMAILS_PER_CATEGORY)]
            rng.shuffle(kinds)
            for layout, kind in zip(layouts, kinds, strict=True):
                group.append({"category": category, "layout": layout, "is_order": False, "kind": kind, "lines": []})
        else:
            sizes = [rng.choice(LINE_COUNTS[category]) for _ in range(EMAILS_PER_CATEGORY)]
            total = sum(sizes)
            traps = []
            for trap, share in TRAP_SHARES:
                traps += [trap] * round(share * total)
            traps += ["plain"] * (total - len(traps))
            rng.shuffle(traps)
            weights_key = (
                "table"
                if category in TABLE_CATEGORIES
                else "conversational"
                if category == "body_conversational"
                else "written"
            )
            for layout, size in zip(layouts, sizes, strict=True):
                email_traps, traps = traps[:size], traps[size:]
                in_catalog = sum(t != "out_of_catalog" for t in email_traps)
                chosen = sorted(products, key=lambda p: (uses[p.sku], rng.random()))[:in_catalog]
                items = rng.sample(OUT_OF_CATALOG_ITEMS, size - in_catalog)
                lines = []
                for trap in email_traps:
                    product = None if trap == "out_of_catalog" else chosen.pop()
                    unit_style, quantity, extra = _unit(rng, weights_key, product)
                    line = {
                        "expected_sku": product.sku if product else None,
                        "expected_quantity": quantity,
                        "unit_style": unit_style,
                        "trap": {"style": trap, **extra},
                    }
                    if product:
                        uses[product.sku] += 1
                        line["sale_unit"] = product.sale_unit
                    else:
                        line["trap"]["requested_item"] = items.pop()
                    if trap == "near_miss":
                        line["trap"]["confusable_with"] = _closest_sibling(product, siblings)
                    lines.append(line)
                group.append({"category": category, "layout": layout, "is_order": True, "lines": lines})
        rng.shuffle(group)
        for i, email in enumerate(group):
            email["split"] = "dev" if i < DEV_PER_CATEGORY else "test"
        emails += group

    # Emails are numbered in a seeded order so ids do not reveal the category.
    rng.shuffle(emails)
    for number, email in enumerate(emails, start=1):
        customer = rng.choice(CUSTOMERS)
        email["id"] = f"EML-{number:04d}"
        email["dataset_version"] = DATASET_VERSION
        email["customer_code"] = customer.code
        email["sender"] = customer.email
        email["source"] = SOURCE[email["category"]]
        email["attachment"] = _attachment_name(email["category"], email["layout"], number)
        email["date"] = format_datetime(BASE_DATE + timedelta(hours=7 * number, minutes=13 * number % 60))
        if email["layout"] == "extra_columns":
            email["cost_centre"] = f"CC-{rng.randint(100, 999)}"
        for n, line in enumerate(email["lines"], start=1):
            line["line_id"] = f"{email['id']}-L{n}"
            line["location"] = email["attachment"] or "body"
    return emails


# ---------- rendering ----------


def _table_rows(planned: dict, text: dict) -> list[tuple[str, int, str]]:
    """Product text, quantity and unit of every table row, taken from the plan."""
    rows = []
    for line in planned["lines"]:
        if line["unit_style"] == "items_to_packs":
            quantity = line["trap"]["items_requested"]
            unit = "pairs" if "pairs" in line["sale_unit"] else "units"
        else:
            quantity = line["expected_quantity"]
            unit = line.get("sale_unit", "unit")
        rows.append((text["lines"][line["line_id"]], quantity, unit))
    return rows


def render_pdf(planned: dict, text: dict) -> bytes:
    from fpdf import FPDF

    customer = CUSTOMER_BY_CODE[planned["customer_code"]]
    number = planned["id"].split("-")[1]
    pdf = FPDF(orientation="L", format="A4")
    pdf.set_creation_date(BASE_DATE)
    pdf.set_title(planned["attachment"])
    pdf.set_author(customer.name)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    if planned["layout"] == "order_form":
        pdf.cell(0, 10, "PURCHASE ORDER", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=10)
        for label in (f"Customer: {customer.name}", f"Order reference: PO-{number}", f"City: {customer.city}"):
            pdf.cell(0, 6, label, new_x="LMARGIN", new_y="NEXT")
        header, widths = ("Product", "Quantity", "Unit"), (170, 30, 50)
        rows = [(p, str(q), u) for p, q, u in _table_rows(planned, text)]
    else:
        pdf.cell(0, 10, "DELIVERY NOTE - REQUESTED ITEMS", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=10)
        for label in (f"Deliver to: {customer.name}, {customer.city}", f"Note number: DN-{number}"):
            pdf.cell(0, 6, label, new_x="LMARGIN", new_y="NEXT")
        header, widths = ("Ref", "Description", "Qty", "Unit"), (15, 160, 25, 50)
        rows = [(str(n), p, str(q), u) for n, (p, q, u) in enumerate(_table_rows(planned, text), start=1)]
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 9)
    for value, width in zip(header, widths, strict=True):
        pdf.cell(width, 7, value, border=1)
    pdf.ln()
    pdf.set_font("Helvetica", size=9)
    for row in rows:
        for value, width in zip(row, widths, strict=True):
            pdf.cell(width, 7, value, border=1)
        pdf.ln()
    return bytes(pdf.output())


def _fixed_zip(data: bytes) -> bytes:
    """Rewrite a zip archive with fixed entry dates so the bytes do not depend on the clock.

    openpyxl stamps `dcterms:modified` with the save time, so it is reset to the creation date.
    """
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + BASE_DATE.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>",
                    content,
                )
            member = zipfile.ZipInfo(info.filename, date_time=FILE_DATE)
            member.create_system = 0  # the default is 3 outside Windows, which changes the bytes
            target.writestr(member, content)
    return out.getvalue()


def render_xlsx(planned: dict, text: dict) -> bytes:
    from openpyxl import Workbook

    customer = CUSTOMER_BY_CODE[planned["customer_code"]]
    workbook = Workbook()
    stamp = BASE_DATE.replace(tzinfo=None)
    workbook.properties.created = stamp
    workbook.properties.modified = stamp
    workbook.properties.creator = customer.name
    sheet = workbook.active
    sheet.title = "Order"
    if planned["layout"] == "header_row":
        sheet.append(["Product", "Quantity", "Unit"])
        for row in _table_rows(planned, text):
            sheet.append(list(row))
    else:
        sheet.append(["Line", "Product", "Quantity", "Unit", "Cost centre", "Requested by"])
        for n, (product, quantity, unit) in enumerate(_table_rows(planned, text), start=1):
            sheet.append([n, product, quantity, unit, planned["cost_centre"], customer.contact_name])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _fixed_zip(buffer.getvalue())


def render_txt(planned: dict, text: dict) -> bytes:
    return ("\n".join(text["lines"][line["line_id"]] for line in planned["lines"]) + "\n").encode("utf-8")


def render_email(planned: dict, text: dict) -> bytes:
    """The `.eml` file of one planned email and its written texts; same input, same bytes."""
    customer = CUSTOMER_BY_CODE[planned["customer_code"]]
    message = EmailMessage()
    message["From"] = f"{customer.contact_name} <{planned['sender']}>"
    message["To"] = SHOP_ADDRESS
    message["Subject"] = text["subject"]
    message["Date"] = planned["date"]
    message["Message-ID"] = f"<{planned['id'].lower()}@dataset.purchase-cycle.example>"
    if planned["layout"] == "html_body":
        paragraphs = "".join(
            f"<p>{'<br>'.join(html.escape(line) for line in block.splitlines())}</p>"
            for block in text["body"].split("\n\n")
        )
        message.set_content(f"<html><body>{paragraphs}</body></html>\n", subtype="html")
    else:
        message.set_content(text["body"] + "\n")
    name = planned["attachment"]
    if name:
        renderer, mime = {
            ".txt": (render_txt, "text/plain"),
            ".pdf": (render_pdf, "application/pdf"),
            ".xlsx": (render_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        }[Path(name).suffix]
        data = renderer(planned, text)
        if mime == "text/plain":
            message.add_attachment(data.decode("utf-8"), filename=name)
        else:
            maintype, subtype = mime.split("/")
            message.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
        message.set_boundary(f"=={planned['id']}-boundary==")
    return bytes(message)


# ---------- validation ----------


def _flat(text: str) -> str:
    """Whitespace-insensitive form used to find a line text inside a rendered text."""
    return " ".join(text.split())


def _contains_name(text: str, sku: str) -> bool:
    return f" {match_key(PRODUCT_BY_SKU[sku].name)} " in f" {match_key(text)} "


def validate_line(planned: dict, line: dict, line_text: str) -> list[str]:
    """Reasons one written line is rejected; empty when it passes."""
    lid, sku, quantity = line["line_id"], line["expected_sku"], line["expected_quantity"]
    errors = []
    if sku is not None and sku not in PRODUCT_BY_SKU:
        errors.append(f"{lid}: expected SKU {sku} does not exist")
    if not isinstance(quantity, int) or isinstance(quantity, bool) or quantity < 1:
        errors.append(f"{lid}: expected quantity {quantity!r} is not a positive whole number")
    if (sku is None) != (line["trap"]["style"] == "out_of_catalog"):
        errors.append(f"{lid}: only out_of_catalog lines expect a null SKU")
    text = line_text.strip()
    if not text:
        return errors + [f"{lid}: line text is empty"]
    if "\n" in text or "\t" in text:
        errors.append(f"{lid}: line text must be one line")
    table = planned["category"] in TABLE_CATEGORIES
    limit = MAX_TABLE_TEXT_CHARS if table else MAX_LINE_CHARS
    if len(text) > limit:
        errors.append(f"{lid}: line text longer than {limit} characters")
    if planned["category"] == "pdf_attachment" and not text.isascii():
        errors.append(f"{lid}: PDF text must be plain ASCII")
    if errors:
        return errors
    style = line["trap"]["style"]
    if style == "out_of_catalog" and match_key(text) in {match_key(p.name) for p in PRODUCTS}:
        errors.append(f"{lid}: out_of_catalog text is a catalog name")
    if style in ("typo", "near_miss") and _contains_name(text, sku):
        errors.append(f"{lid}: {style} text contains the exact catalog name")
    if style == "near_miss" and _contains_name(text, line["trap"]["confusable_with"]):
        errors.append(f"{lid}: near_miss text names the confusable product")
    if not table:  # the customer writes the quantity in the line text
        unit = line["unit_style"]
        words = text.lower()
        if unit == "sale_units" and not has_number(text, quantity):
            errors.append(f"{lid}: quantity {quantity} not written as a figure")
        elif unit == "items_to_packs" and not has_number(text, line["trap"]["items_requested"]):
            errors.append(f"{lid}: item count {line['trap']['items_requested']} not written as a figure")
        elif unit == "quantity_words" and (has_number(text, quantity) or not re.search(r"[a-z]", words)):
            errors.append(f"{lid}: quantity {quantity} must be written in words")
        elif unit == "dozen" and "dozen" not in words:
            errors.append(f"{lid}: dozen expression missing")
        elif unit == "half_dozen" and not ("half" in words and "dozen" in words):
            errors.append(f"{lid}: half dozen expression missing")
        elif unit == "couple" and "couple" not in words:
            errors.append(f"{lid}: couple expression missing")
    return errors


def validate_email(planned: dict, text: dict) -> list[str]:
    """Reasons a written email is rejected before rendering; empty when it passes."""
    errors = []
    subject, body = (text.get("subject") or "").strip(), (text.get("body") or "").strip()
    written = text.get("lines") or {}
    if not subject or "\n" in subject or len(subject) > MAX_SUBJECT_CHARS:
        errors.append(f"subject must be one line of 1 to {MAX_SUBJECT_CHARS} characters")
    if not body or len(body) > MAX_BODY_CHARS:
        errors.append(f"body must have 1 to {MAX_BODY_CHARS} characters")
    planned_ids = [line["line_id"] for line in planned["lines"]]
    if sorted(written) != sorted(planned_ids):
        errors.append(f"written lines {sorted(written)} do not match the planned lines {planned_ids}")
        return errors
    if planned["is_order"] != bool(planned["lines"]):
        errors.append("an order needs lines and a non-order has none")
    for line in planned["lines"]:
        errors += validate_line(planned, line, written[line["line_id"]])
    if errors:
        return errors
    category, flat_body = planned["category"], _flat(body).lower()
    texts = [_flat(written[i]).lower() for i in planned_ids]
    if category in BODY_CATEGORIES:
        missing = [i for i, t in zip(planned_ids, texts, strict=True) if t not in flat_body]
        if missing:
            errors.append(f"line texts missing from the body: {', '.join(missing)}")
    elif category != "not_an_order":
        if len(body) > MAX_SHORT_BODY_CHARS:
            errors.append(f"{category} body longer than {MAX_SHORT_BODY_CHARS} characters")
        if any(t in flat_body for t in texts):
            errors.append(f"{category} body repeats a line text; the lines belong in the attachment")
    if category == "body_list":
        body_lines = [_flat(b).lower() for b in body.splitlines()]
        shared = [
            i
            for i, t in zip(planned_ids, texts, strict=True)
            if not any(t in b and sum(o in b for o in texts) == 1 for b in body_lines)
        ]
        if shared:
            errors.append(f"body_list lines must each sit on their own body line: {', '.join(shared)}")
    if category == "body_conversational":
        rest = flat_body
        for t in texts:
            rest = rest.replace(t, " ")
        if len(_flat(rest)) < MIN_CONTEXT_CHARS:
            errors.append(f"body_conversational needs at least {MIN_CONTEXT_CHARS} characters of context")
    return errors


def check_rendered(planned: dict, text: dict, data: bytes) -> list[str]:
    """Parse the rendered file with the batch A intake and confirm every planned line is in its location."""
    try:
        email = parse_email(data)
    except EmailRejected as error:
        return [f"rendered email cannot be read: {error}"]
    errors = []
    if email["sender"] != planned["sender"]:
        errors.append(f"rendered sender {email['sender']} is not {planned['sender']}")
    if email["ignored"]:
        errors.append(f"rendered attachments ignored: {email['ignored']}")
    located = {"body": _flat(email["body"]).lower()}
    located.update({a["name"]: _flat(a["text"]).lower() for a in email["attachments"]})
    for line in planned["lines"]:
        where = located.get(line["location"])
        line_text = _flat(text["lines"][line["line_id"]]).lower()
        if where is None or line_text not in where:
            errors.append(f"{line['line_id']}: planned line missing from the rendered {line['location']}")
    return errors


def email_key(planned: dict, text: dict) -> str:
    """Normalised text of one email, used to reject duplicates."""
    lines = [text["lines"][line["line_id"]] for line in planned["lines"]]
    return match_key(" ".join([text.get("body") or "", *lines]))


def load_texts(directory: Path = TEXTS_DIR) -> dict[str, dict]:
    texts = {}
    for path in sorted(directory.glob("*.jsonl")):
        for row in read_jsonl(path):
            texts[row["id"]] = row
    return texts


def build_dataset(plan: list[dict], texts: dict[str, dict]) -> tuple[list[dict], dict[str, bytes], dict[str, list]]:
    """Validate and render every planned email; return dataset rows, rendered files and rejected ids with reasons."""
    rows, files, rejected, seen = [], {}, {}, {}
    for planned in plan:
        text = texts.get(planned["id"]) or {}
        errors = validate_email(planned, text)
        if not errors:
            key = email_key(planned, text)
            if key in seen:
                errors.append(f"duplicates {seen[key]} after normalisation")
            seen.setdefault(key, planned["id"])
        if not errors:
            data = render_email(planned, text)
            errors = check_rendered(planned, text, data)
            files[planned["id"]] = data
        if errors:
            rejected[planned["id"]] = errors
            continue
        lines = [dict(line, text=text["lines"][line["line_id"]]) for line in planned["lines"]]
        rows.append(
            dict(planned, subject=text["subject"], body=text["body"], lines=lines, file=f"emails/{planned['id']}.eml")
        )
    return rows, files, rejected


def load_dataset(path: Path = DATASET_PATH) -> list[dict]:
    return read_jsonl(path)


# ---------- second-pass review ----------


def _line_similarity(label_text: str, read_text: str) -> float:
    a, b = match_key(label_text), match_key(read_text)
    if a and b and (a in b or b in a):
        return 1.0
    return SequenceMatcher(None, a, b).ratio()


def review_disagreements(email: dict, review: dict) -> list[dict]:
    """Differences between the labels of one email and a blind annotator's reading of it.

    Each labelled line is paired with the read line of the same SKU, else with
    the most similar read text (at least MIN_LINE_SIMILARITY); unpaired lines on
    either side are disagreements too.
    """
    found = []
    if review["is_order"] != email["is_order"]:
        found.append({"field": "is_order", "label": email["is_order"], "read": review["is_order"]})
    unpaired = list(review["lines"])
    for line in email["lines"]:
        pair = next((r for r in unpaired if r["sku"] and r["sku"] == line["expected_sku"]), None)
        if pair is None and unpaired:
            score, pair = max(((_line_similarity(line["text"], r["text"]), r) for r in unpaired), key=lambda t: t[0])
            if score < MIN_LINE_SIMILARITY:
                pair = None
        if pair is None:
            found.append({"field": "line", "line_id": line["line_id"], "label": line["expected_sku"], "read": None})
            continue
        unpaired.remove(pair)
        for field, label_key, read_key in (
            ("sku", "expected_sku", "sku"),
            ("quantity", "expected_quantity", "quantity"),
        ):
            if pair[read_key] != line[label_key]:
                found.append(
                    {"field": field, "line_id": line["line_id"], "label": line[label_key], "read": pair[read_key]}
                )
    found += [{"field": "line", "line_id": None, "label": None, "read": r["text"]} for r in unpaired]
    return found


def build_review(dataset: list[dict], reviews: list[dict], previous: list[dict]) -> list[dict]:
    """One record per email: the blind reading, its disagreements with the labels and the decision.

    Decisions and justifications of an earlier review file are kept; a new
    disagreement starts as `pending` until it is fixed or justified by hand.
    """
    by_id = {r["id"]: r for r in reviews}
    kept = {r["id"]: r for r in previous}
    records = []
    for email in dataset:
        review = by_id[email["id"]]
        found = review_disagreements(email, review)
        old = kept.get(email["id"], {})
        decision = old.get("decision", "pending") if found else "agree"
        records.append(
            {
                "id": email["id"],
                "read": {"is_order": review["is_order"], "lines": review["lines"]},
                "notes": review.get("notes", ""),
                "disagreements": found,
                "decision": decision,
                "justification": old.get("justification", "") if found else "",
            }
        )
    return records


def cmd_review(args) -> int:
    reviews = [row for path in args.annotations for row in read_jsonl(Path(path))]
    previous = read_jsonl(REVIEW_PATH) if REVIEW_PATH.exists() else []
    records = build_review(load_dataset(), reviews, previous)
    write_jsonl(REVIEW_PATH, records)
    counts = Counter(r["decision"] for r in records)
    print(f"Reviewed {len(records)} emails -> {REVIEW_PATH}")
    print("  " + ", ".join(f"{d} {counts[d]}" for d in (*REVIEW_DECISIONS, "pending")))
    return 1 if counts["pending"] else 0


# ---------- owner audit ----------


def _expected_lines(email: dict) -> str:
    out = []
    for line in email["lines"]:
        product = PRODUCT_BY_SKU.get(line["expected_sku"])
        target = f"{product.sku} {product.name}" if product else "NOT IN CATALOG"
        unit = f" ({line['sale_unit']})" if product else ""
        out.append(f"{line['line_id']}: {line['text']} -> {target} x {line['expected_quantity']}{unit}")
    return "\n".join(out)


def create_audit(path: Path = AUDIT_PATH, dataset_path: Path = DATASET_PATH) -> int:
    if path.exists():
        print(f"{path} already exists; it may hold the owner's verdicts, so it is not overwritten")
        return 1
    dataset = load_dataset(dataset_path)
    sample = sorted(random.Random(AUDIT_SEED).sample(dataset, AUDIT_SIZE), key=lambda e: e["id"])
    path.parent.mkdir(parents=True, exist_ok=True)
    # Semicolons and a BOM so a Spanish-locale Excel opens the columns directly, as in phases 01 and 02.
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=AUDIT_COLUMNS, delimiter=";")
        writer.writeheader()
        for email in sample:
            data = (dataset_path.parent / email["file"]).read_bytes()
            writer.writerow(
                {
                    "id": email["id"],
                    "category": email["category"],
                    "file": email["file"],
                    "email_text": model_text(parse_email(data)),
                    "is_order": "yes" if email["is_order"] else "no",
                    "expected_lines": _expected_lines(email),
                    "verdict": "",
                    "comment": "",
                }
            )
    print(f"Wrote {len(sample)} emails -> {path}")
    return 0


def report_audit(path: Path = AUDIT_PATH) -> int:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh, delimiter=";"))
    pending = [r["id"] for r in rows if r["verdict"].strip().lower() not in VERDICTS]
    if pending:
        print(f"{len(pending)} rows still without a verdict (ok or wrong), first: {', '.join(pending[:5])}")
        return 2
    wrong = [r for r in rows if r["verdict"].strip().lower() == "wrong"]
    low, high = wilson_interval(len(wrong), len(rows))
    print(
        f"email_order_extraction audit n={len(rows)} wrong labels={len(wrong)} "
        f"error rate={len(wrong) / len(rows):.1%}  95% CI [{low:.1%}, {high:.1%}]"
    )
    for r in wrong:
        print(f"  {r['id']}: {r['comment'] or 'no comment'}")
    if len(wrong) > MAX_WRONG:
        print(f"FAILED: more than {MAX_WRONG} wrong label in the audit sample")
        return 1
    return 0


# ---------- commands ----------


def cmd_plan(args) -> int:
    plan = build_plan(args.seed)
    write_jsonl(PLAN_PATH, plan)
    lines = sum(len(e["lines"]) for e in plan)
    print(f"Planned {len(plan)} emails with {lines} expected lines, seed {args.seed} -> {PLAN_PATH}")
    return 0


def cmd_check(args) -> int:
    plan = {e["id"]: e for e in read_jsonl(PLAN_PATH)}
    batch = read_jsonl(Path(args.batch))
    others = {k: v for k, v in load_texts().items() if k not in {row["id"] for row in batch}}
    subset = [plan[row["id"]] for row in batch]
    known = [e for e in plan.values() if e["id"] in others]
    _, _, rejected = build_dataset(known + subset, {**others, **{row["id"]: row for row in batch}})
    failed = 0
    for planned in subset:
        if planned["id"] in rejected:
            failed += 1
            print(f"{planned['id']}: {'; '.join(rejected[planned['id']])}")
    print(f"{len(batch)} emails checked, {failed} rejected")
    return 1 if failed else 0


def cmd_build(args) -> int:
    plan = read_jsonl(PLAN_PATH)
    rows, files, rejected = build_dataset(plan, load_texts())
    for email_id, errors in rejected.items():
        print(f"{email_id}: {'; '.join(errors)}")
    if rejected:
        print(f"{len(rejected)} of {len(plan)} emails rejected; rewrite their texts and build again")
        return 1
    EMAILS_DIR.mkdir(parents=True, exist_ok=True)
    for email_id, data in files.items():
        (EMAILS_DIR / f"{email_id}.eml").write_bytes(data)
    write_jsonl(DATASET_PATH, rows)
    counts = Counter(row["category"] for row in rows)
    lines = sum(len(row["lines"]) for row in rows)
    print(f"Built {len(rows)} emails with {lines} expected lines -> {DATASET_PATH}")
    print("  " + ", ".join(f"{c} {counts[c]}" for c in CATEGORIES))
    return 0


def add_commands(sub) -> None:
    dataset = sub.add_parser("email-dataset", help="plan, check and build the email order extraction dataset")
    dsub = dataset.add_subparsers(dest="email_dataset_command", required=True)
    plan = dsub.add_parser("plan", help="write the seeded email plan")
    plan.add_argument("--seed", type=int, default=SEED)
    plan.set_defaults(handler=cmd_plan)
    check = dsub.add_parser("check", help="validate and render a batch of written emails against the plan")
    check.add_argument("batch")
    check.set_defaults(handler=cmd_check)
    dsub.add_parser("build", help="validate, render the .eml files and write the dataset").set_defaults(
        handler=cmd_build
    )
    review = dsub.add_parser("review", help="compare blind second-pass annotations with the labels")
    review.add_argument("annotations", nargs="+", help="JSONL files with one blind reading per email")
    review.set_defaults(handler=cmd_review)
    audit = sub.add_parser("email-audit", help="owner audit of the email order extraction labels")
    asub = audit.add_subparsers(dest="email_audit_command", required=True)
    asub.add_parser("create", help="write the review file").set_defaults(handler=lambda args: create_audit())
    asub.add_parser("report", help="compute the label error rate").set_defaults(handler=lambda args: report_audit())
